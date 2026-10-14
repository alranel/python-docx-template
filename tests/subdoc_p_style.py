# -*- coding: utf-8 -*-
"""Tests for {{p var }} paragraph style inheritance onto Subdoc bodies."""
import io
import os
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docxtpl import DocxTemplate, scan_p_placeholder_styles

os.chdir(os.path.dirname(__file__))
os.makedirs("output", exist_ok=True)


def _template_with_p_placeholder(varname, p_style=None, alignment=None):
    doc = Document()
    p = doc.add_paragraph()
    if p_style:
        try:
            p.style = p_style
        except KeyError:
            pPr = p._element.get_or_add_pPr()
            pStyle = pPr.find(qn("w:pStyle"))
            if pStyle is None:
                pStyle = OxmlElement("w:pStyle")
                pPr.insert(0, pStyle)
            pStyle.set(qn("w:val"), p_style)
    if alignment is not None:
        p.alignment = alignment
    p.add_run("{{p %s }}" % varname)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _etree_tostring_body(subdoc):
    from lxml import etree

    return etree.tostring(subdoc.element.body, encoding="unicode")


class SubdocPStyleTests(unittest.TestCase):
    def test_scan_p_placeholder_styles(self):
        tpl = DocxTemplate(
            _template_with_p_placeholder(
                "notes", p_style="Quote", alignment=WD_ALIGN_PARAGRAPH.CENTER
            )
        )
        styles = scan_p_placeholder_styles(tpl)
        self.assertIn("notes", styles)
        self.assertEqual(styles["notes"].p_style, "Quote")
        self.assertEqual(styles["notes"].alignment, "center")

    def test_subdoc_inherits_placeholder_style_on_render(self):
        tpl = DocxTemplate(
            _template_with_p_placeholder("notes", p_style="Quote")
        )
        sd = tpl.new_subdoc()
        sd.add_paragraph("Styled body")
        tpl.render({"notes": sd})
        tpl.save("output/subdoc_p_style.docx")

        styled = False
        for p in sd.paragraphs:
            if p.text.strip() == "Styled body":
                if p.style and p.style.name == "Quote":
                    styled = True
        if not styled:
            xml = _etree_tostring_body(sd)
            self.assertTrue(
                "Quote" in xml,
                "expected Quote style on Subdoc body after render",
            )

    def test_heading_paragraphs_not_overwritten(self):
        tpl = DocxTemplate(
            _template_with_p_placeholder("notes", p_style="Quote")
        )
        sd = tpl.new_subdoc()
        heading = sd.add_paragraph("Keep heading")
        try:
            heading.style = "Heading 1"
        except KeyError:
            self.skipTest("Heading 1 style unavailable")
        plain = sd.add_paragraph("Plain body")
        tpl.render({"notes": sd})

        self.assertEqual(heading.style.name, "Heading 1")
        if plain.style and plain.style.name == "Quote":
            return
        xml = _etree_tostring_body(sd)
        self.assertIn("Quote", xml)


if __name__ == "__main__":
    unittest.main()
