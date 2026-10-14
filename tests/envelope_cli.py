# -*- coding: utf-8 -*-
"""Tests for JSON CLI envelopes (subdoc + richtext)."""
import base64
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document

from docxtpl import DocxTemplate, RichText, expand_context

os.chdir(os.path.dirname(__file__))
os.makedirs("output", exist_ok=True)


def _docx_bytes_with_paragraph(text):
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _template_with_p_placeholder(varname):
    """Build a minimal .docx whose body is a single {{p varname }} paragraph."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{p %s }}" % varname)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _template_with_r_placeholder(varname, prefix="Label: "):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run(prefix)
    p.add_run("{{r %s }}" % varname)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


class EnvelopeTests(unittest.TestCase):
    def test_expand_richtext_envelope(self):
        tpl = DocxTemplate(_template_with_r_placeholder("body"))
        ctx = expand_context(
            tpl,
            {
                "body": {
                    "__docxtpl__": "richtext",
                    "runs": [
                        {"text": "Hello ", "bold": False},
                        {"text": "world", "bold": True},
                    ],
                }
            },
        )
        self.assertIsInstance(ctx["body"], RichText)
        tpl.render(ctx)
        out = "output/envelope_richtext.docx"
        tpl.save(out)
        self.assertTrue(os.path.isfile(out))

    def test_expand_subdoc_envelope(self):
        mini = _docx_bytes_with_paragraph("Inserted block")
        tpl = DocxTemplate(_template_with_p_placeholder("body"))
        ctx = expand_context(
            tpl,
            {
                "body": {
                    "__docxtpl__": "subdoc",
                    "docx_base64": base64.b64encode(mini).decode("ascii"),
                }
            },
        )
        from docxtpl.subdoc import Subdoc

        self.assertIsInstance(ctx["body"], Subdoc)
        tpl.render(ctx)
        out = "output/envelope_subdoc.docx"
        tpl.save(out)
        self.assertTrue(os.path.isfile(out))
        with zipfile.ZipFile(out) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        self.assertIn("Inserted block", xml)

    def test_expand_subdoc_path_envelope(self):
        from docxtpl.subdoc import Subdoc

        tpl = DocxTemplate(_template_with_p_placeholder("body"))
        ctx = expand_context(
            tpl,
            {
                "body": {
                    "__docxtpl__": "subdoc",
                    "path": "templates/envelope_cli_sub.docx",
                }
            },
        )
        self.assertIsInstance(ctx["body"], Subdoc)
        tpl.render(ctx)
        out = "output/envelope_subdoc_path.docx"
        tpl.save(out)
        with zipfile.ZipFile(out) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        self.assertIn("Inserted via CLI", xml)

    def test_cli_end_to_end_envelopes(self):
        out = "output/envelope_cli_e2e.docx"
        if os.path.exists(out):
            os.unlink(out)
        result = subprocess.run(
            [
                sys.executable,
                "-m",
                "docxtpl",
                "templates/envelope_cli_tpl.docx",
                "templates/envelope_cli.json",
                out,
                "-o",
                "-q",
            ],
            capture_output=True,
            text=True,
        )
        self.assertEqual(
            result.returncode,
            0,
            "CLI failed: stdout=%r stderr=%r" % (result.stdout, result.stderr),
        )
        self.assertTrue(os.path.isfile(out))
        with zipfile.ZipFile(out) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        self.assertIn("Hello", xml)
        self.assertIn("world", xml)
        self.assertIn("Inserted via CLI", xml)

    def test_cli_end_to_end_base64_subdoc(self):
        mini = _docx_bytes_with_paragraph("Base64 block")
        with tempfile.TemporaryDirectory() as tmp:
            tpl_path = os.path.join(tmp, "tpl.docx")
            json_path = os.path.join(tmp, "ctx.json")
            out_path = os.path.join(tmp, "out.docx")
            buf = _template_with_p_placeholder("body")
            with open(tpl_path, "wb") as fh:
                fh.write(buf.read())
            with open(json_path, "w") as fh:
                json.dump(
                    {
                        "body": {
                            "__docxtpl__": "subdoc",
                            "docx_base64": base64.b64encode(mini).decode("ascii"),
                        }
                    },
                    fh,
                )
            result = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "docxtpl",
                    tpl_path,
                    json_path,
                    out_path,
                    "-o",
                    "-q",
                ],
                capture_output=True,
                text=True,
            )
            self.assertEqual(
                result.returncode,
                0,
                "CLI failed: stdout=%r stderr=%r" % (result.stdout, result.stderr),
            )
            with zipfile.ZipFile(out_path) as zf:
                xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("Base64 block", xml)

    def test_unknown_envelope_type(self):
        tpl = DocxTemplate(_template_with_r_placeholder("body"))
        with self.assertRaises(RuntimeError):
            expand_context(
                tpl, {"body": {"__docxtpl__": "nope", "runs": []}}
            )

    def test_nested_list_expansion(self):
        tpl = DocxTemplate(_template_with_r_placeholder("body"))
        ctx = expand_context(
            tpl,
            {
                "rows": [
                    {
                        "body": {
                            "__docxtpl__": "richtext",
                            "runs": [{"text": "row1", "bold": True}],
                        }
                    }
                ]
            },
        )
        self.assertIsInstance(ctx["rows"][0]["body"], RichText)


if __name__ == "__main__":
    unittest.main()
